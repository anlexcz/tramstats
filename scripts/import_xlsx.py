"""Usage: python scripts/import_xlsx.py /path/to/export.xlsx (openpyxl required)."""
import openpyxl, datetime, json, re, sys
from pathlib import Path
root=Path(__file__).resolve().parents[1]; out=root/'data'; out.mkdir(exist_ok=True)
f=openpyxl.load_workbook(sys.argv[1],data_only=False); v=openpyxl.load_workbook(sys.argv[1],data_only=True)
def clean(x):
 if x is None or x=='': return None
 if isinstance(x,(int,float)): return str(int(x)) if int(x)==x else str(x)
 return str(x)
def time(x):
 if isinstance(x,(datetime.time,datetime.datetime)): return x.strftime('%H:%M')
 if isinstance(x,datetime.timedelta):
  m=round(x.total_seconds()/60); return f'{m//60}:{m%60:02}'
 if isinstance(x,(int,float)):
  m=round(x*1440); return f'{m//60:02}:{m%60:02}'
 return clean(x)
def dump(name,data): (out/name).write_text(json.dumps(data,ensure_ascii=False,separators=(',',':'))+'\n')
manifest={'schemaVersion':1,'source':'https://docs.google.com/spreadsheets/d/1G-WI9H1T7sSH5g6IBZXzbCx-3lbVZ7Df8mrpM6jxR6E/edit','exportedAt':'2026-10-29','years':[]}
for year in range(2020,2027):
 s=f[str(year)]; cached=v[str(year)]; assert [c.value.lower() for c in s[1][:17]][:17]==[x.lower() for x in ['Datum','Den','PS','SS','Č','Vozovna','K','L','P','Typ','Přední vůz','Zadní vůz','Od','Do','Délka','km','Poznámka']]
 shifts=[]; current=None
 for idx,row in enumerate(s.iter_rows(min_row=2),2):
  a=[c.value for c in row[:17]]; cv=[c.value for c in cached[idx][:17]]
  if isinstance(a[0],datetime.datetime) and (current is None or current['date'] != a[0].date().isoformat()):
   current={'date':a[0].date().isoformat(),'day':clean(cv[1]),'plannedShift':clean(a[2]),'actualShift':clean(a[3]),'depot':clean(a[5]),'parts':[]};shifts.append(current)
  elif not any(a[i] is not None for i in [4,5,6,7,8,9,10,11,12,13,15,16]): continue
  if current is None: raise ValueError(f'Orphan row {year}/{idx}')
  if not any(a[i] is not None for i in [4,6,7,8,9,10,11,12,13,15,16]): continue
  keys=['partNo','depot','course','line','order','vehicleType','frontVehicle','rearVehicle'];p=dict(zip(keys,[clean(x) for x in a[4:12]]))
  p.update(start=time(a[12]),end=time(a[13]),duration=time(cv[14]),km=cv[15] if isinstance(cv[15],(int,float)) else None,kmFormula=a[15] if isinstance(a[15],str) and a[15].startswith('=') else None,note=clean(a[16]),sourceRow=idx)
  if p['km'] is None and isinstance(a[15],(int,float)):p['km']=a[15]
  current['parts'].append(p)
 bp=[];bs=f[f'{year}_BP'];bc=v[f'{year}_BP'];assert [c.value for c in bs[1]][:7]==['Datum','Od','Do','Délka','Od místa','Do místa','Poznámka']
 for idx,row in enumerate(bs.iter_rows(min_row=2),2):
  a=[c.value for c in row[:7]]
  if not isinstance(a[0],datetime.datetime):continue
  bp.append(dict(zip(['date','start','end','duration','fromPlace','toPlace','note'],[a[0].date().isoformat(),time(a[1]),time(a[2]),time(bc.cell(idx,4).value),clean(a[4]),clean(a[5]),clean(a[6])])))
 assert len({s['date'] for s in shifts})==len(shifts)
 dump(f'shifts-{year}.json',shifts);dump(f'breaks-{year}.json',bp)
 item={'year':year,'shifts':f'shifts-{year}.json','breaks':f'breaks-{year}.json','days':len(shifts),'parts':sum(len(x['parts']) for x in shifts),'breakCount':len(bp),'km':round(sum(p['km'] or 0 for x in shifts for p in x['parts']),3),'from':min(x['date'] for x in shifts),'to':max(x['date'] for x in shifts)};manifest['years'].append(item);print(item)
dump('manifest.json',manifest)
